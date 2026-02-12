"""
PPT Service for generating presentations
"""

import json
import re
from typing import Dict, Any, List, Optional
from ..api.models import PPTGenerationRequest, PPTOutline
import docx
import PyPDF2
import io

class PPTService:
    """Service for PPT generation and processing"""
    
    def __init__(self):
        self.scenario_configs = {
            "general": {
                "color_scheme": "#2E86AB",
                "font_family": "Arial, sans-serif",
                "style_class": "general-theme"
            },
            "tourism": {
                "color_scheme": "#27AE60",
                "font_family": "Georgia, serif",
                "style_class": "tourism-theme"
            },
            "education": {
                "color_scheme": "#E74C3C",
                "font_family": "Comic Sans MS, cursive",
                "style_class": "education-theme"
            },
            "analysis": {
                "color_scheme": "#34495E",
                "font_family": "Helvetica, sans-serif",
                "style_class": "analysis-theme"
            },
            "history": {
                "color_scheme": "#8B4513",
                "font_family": "Times New Roman, serif",
                "style_class": "history-theme"
            },
            "technology": {
                "color_scheme": "#9B59B6",
                "font_family": "Roboto, sans-serif",
                "style_class": "technology-theme"
            },
            "business": {
                "color_scheme": "#1F4E79",
                "font_family": "Calibri, sans-serif",
                "style_class": "business-theme"
            }
        }
    
    async def generate_ppt(self, task_id: str, request: PPTGenerationRequest) -> Dict[str, Any]:
        """Generate complete PPT based on request"""
        try:
            # Step 1: Generate outline
            outline = await self.generate_outline(request)
            
            # Step 2: Generate slides HTML
            slides_html = await self.generate_slides_from_outline(outline, request.scenario)
            
            return {
                "success": True,
                "task_id": task_id,
                "outline": outline,
                "slides_html": slides_html
            }
            
        except Exception as e:
            return {
                "success": False,
                "task_id": task_id,
                "error": str(e)
            }
    
    async def generate_outline(self, request: PPTGenerationRequest) -> PPTOutline:
        """Generate PPT outline based on request"""
        topic = request.topic
        scenario = request.scenario
        language = request.language
        ppt_style = request.ppt_style
        custom_style_prompt = request.custom_style_prompt
        description = request.description
        # Generate slides based on scenario
        slides = []
        
        # Title slide
        slides.append({
            "id": 1,
            "type": "title",
            "title": topic,
            "subtitle": "专业演示" if language == "zh" else "Professional Presentation",
            "content": ""
        })
        
        # Agenda slide
        slides.append({
            "id": 2,
            "type": "agenda",
            "title": "目录" if language == "zh" else "Agenda",
            "subtitle": "",
            "content": self._generate_agenda_content(scenario, language,)
        })
        
        # Content slides based on scenario
        content_slides = self._generate_content_slides(topic, scenario, language)
        slides.extend(content_slides)
        
        # Thank you slide
        slides.append({
            "id": len(slides) + 1,
            "type": "thankyou",
            "title": "谢谢" if language == "zh" else "Thank You",
            "subtitle": "感谢聆听" if language == "zh" else "Thank you for your attention",
            "content": ""
        })
        
        return PPTOutline(
            title=topic,
            slides=slides,
            metadata={
                "scenario": scenario,
                "language": language,
                "total_slides": len(slides),
                "generated_at": "2024-01-01T00:00:00Z"
            }
        )
    
    async def generate_slides_from_outline(self, outline: PPTOutline, scenario: str) -> str:
        """Generate HTML slides from outline"""
        config = self.scenario_configs.get(scenario, self.scenario_configs["general"])
        
        # Generate CSS styles
        css_styles = self._generate_css_styles(config)
        
        # Generate HTML for each slide
        slides_html = []
        for slide in outline.slides:
            slide_html = self._generate_slide_html(slide, config)
            slides_html.append(slide_html)
        
        # Combine into complete HTML document
        complete_html = f"""
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{outline.title}</title>
    <style>
        {css_styles}
    </style>
</head>
<body>
    <div class="presentation-container">
        <div class="slides-wrapper">
            {''.join(slides_html)}
        </div>
        <div class="navigation">
            <button id="prevBtn" onclick="changeSlide(-1)">‹ 上一页</button>
            <span id="slideCounter">1 / {len(outline.slides)}</span>
            <button id="nextBtn" onclick="changeSlide(1)">下一页 ›</button>
        </div>
    </div>
    
    <script>
        let currentSlide = 0;
        const totalSlides = {len(outline.slides)};

        function showSlide(n) {{
            const slides = document.querySelectorAll('.slide');

            // Check if slides exist
            if (!slides || slides.length === 0) {{
                console.error('No slides found');
                return;
            }}

            // Validate slide index
            if (n >= totalSlides) currentSlide = 0;
            if (n < 0) currentSlide = totalSlides - 1;

            // Hide all slides
            slides.forEach(slide => {{
                if (slide && slide.style) {{
                    slide.style.display = 'none';
                }}
            }});

            // Show current slide
            if (slides[currentSlide] && slides[currentSlide].style) {{
                slides[currentSlide].style.display = 'block';
            }}

            // Update counter
            const counter = document.getElementById('slideCounter');
            if (counter) {{
                counter.textContent = `${{currentSlide + 1}} / ${{totalSlides}}`;
            }}
        }}

        function changeSlide(n) {{
            currentSlide += n;
            showSlide(currentSlide);
        }}

        // Keyboard navigation
        document.addEventListener('keydown', function(e) {{
            if (e.key === 'ArrowLeft') changeSlide(-1);
            if (e.key === 'ArrowRight') changeSlide(1);
        }});

        // Initialize when DOM is ready
        document.addEventListener('DOMContentLoaded', function() {{
            showSlide(0);
        }});

        // Fallback initialization
        window.addEventListener('load', function() {{
            setTimeout(() => showSlide(0), 100);
        }});
    </script>
</body>
</html>
        """
        
        return complete_html
    
    async def process_uploaded_file(self, filename: str, content: bytes, file_type: str) -> str:
        """Process uploaded file and extract content"""
        try:
            if file_type == ".docx":
                return self._process_docx(content)
            elif file_type == ".pdf":
                return self._process_pdf(content)
            elif file_type in [".txt", ".md"]:
                return content.decode('utf-8')
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
                
        except Exception as e:
            raise Exception(f"Error processing file: {str(e)}")
    
    def _process_docx(self, content: bytes) -> str:
        """Process DOCX file and extract text"""
        doc = docx.Document(io.BytesIO(content))
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text.strip())
        
        return "\n".join(text_content)
    
    def _process_pdf(self, content: bytes) -> str:
        """Process PDF file and extract text"""
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
        text_content = []
        
        for page in pdf_reader.pages:
            text = page.extract_text()
            if text.strip():
                text_content.append(text.strip())
        
        return "\n".join(text_content)
    
    def _generate_agenda_content(self, scenario: str, language: str) -> str:
        """Generate agenda content based on scenario"""
        agenda_templates = {
            "general": {
                "zh": ["引言", "主要内容", "案例分析", "总结"],
                "en": ["Introduction", "Main Content", "Case Study", "Conclusion"]
            },
            "tourism": {
                "zh": ["目的地概览", "主要景点", "行程安排", "实用信息"],
                "en": ["Destination Overview", "Main Attractions", "Itinerary", "Practical Information"]
            },
            "education": {
                "zh": ["学习目标", "核心概念", "实例说明", "互动活动", "总结回顾"],
                "en": ["Learning Objectives", "Key Concepts", "Examples", "Activities", "Summary"]
            },
            "analysis": {
                "zh": ["问题陈述", "研究方法", "研究发现", "深入分析", "建议方案"],
                "en": ["Problem Statement", "Methodology", "Findings", "Analysis", "Recommendations"]
            },
            "history": {
                "zh": ["历史背景", "时间线", "关键事件", "重要意义", "历史影响"],
                "en": ["Background", "Timeline", "Key Events", "Significance", "Legacy"]
            },
            "technology": {
                "zh": ["技术概览", "核心功能", "优势特点", "实施方案", "未来展望"],
                "en": ["Overview", "Features", "Benefits", "Implementation", "Future"]
            },
            "business": {
                "zh": ["执行摘要", "问题分析", "解决方案", "市场分析", "财务预测"],
                "en": ["Executive Summary", "Problem", "Solution", "Market Analysis", "Financial Projections"]
            }
        }
        
        template = agenda_templates.get(scenario, agenda_templates["general"])
        items = template.get(language, template["en"])
        
        return "\n".join([f"• {item}" for item in items])
    
    def _generate_content_slides(self, topic: str, scenario: str, language: str) -> List[Dict[str, Any]]:
        """Generate content slides based on topic and scenario"""
        slides = []
        
        # Get agenda items to create content slides
        agenda_templates = {
            "general": {
                "zh": ["引言", "主要内容", "案例分析", "总结"],
                "en": ["Introduction", "Main Content", "Case Study", "Conclusion"]
            },
            "tourism": {
                "zh": ["目的地概览", "主要景点", "行程安排", "实用信息"],
                "en": ["Destination Overview", "Main Attractions", "Itinerary", "Practical Information"]
            },
            "education": {
                "zh": ["学习目标", "核心概念", "实例说明", "互动活动", "总结回顾"],
                "en": ["Learning Objectives", "Key Concepts", "Examples", "Activities", "Summary"]
            },
            "analysis": {
                "zh": ["问题陈述", "研究方法", "研究发现", "深入分析", "建议方案"],
                "en": ["Problem Statement", "Methodology", "Findings", "Analysis", "Recommendations"]
            },
            "history": {
                "zh": ["历史背景", "时间线", "关键事件", "重要意义", "历史影响"],
                "en": ["Background", "Timeline", "Key Events", "Significance", "Legacy"]
            },
            "technology": {
                "zh": ["技术概览", "核心功能", "优势特点", "实施方案", "未来展望"],
                "en": ["Overview", "Features", "Benefits", "Implementation", "Future"]
            },
            "business": {
                "zh": ["执行摘要", "问题分析", "解决方案", "市场分析", "财务预测"],
                "en": ["Executive Summary", "Problem", "Solution", "Market Analysis", "Financial Projections"]
            }
        }
        
        template = agenda_templates.get(scenario, agenda_templates["general"])
        items = template.get(language, template["en"])
        
        for i, item in enumerate(items, 3):  # Start from slide 3 (after title and agenda)
            content = self._generate_slide_content(topic, item, scenario, language)
            slides.append({
                "id": i,
                "type": "content",
                "title": item,
                "subtitle": "",
                "content": content
            })
        
        return slides
    
    def _generate_slide_content(self, topic: str, section: str, scenario: str, language: str) -> str:
        """Generate content for a specific slide section"""
        # This is a simplified content generation
        
        if language == "zh":
            content_templates = {
                "引言": f"• {topic}的重要性\n• 本次演示的目标\n• 主要讨论内容概览",
                "主要内容": f"• {topic}的核心要点\n• 关键特征和优势\n• 实际应用场景",
                "案例分析": f"• 成功案例展示\n• 实施过程分析\n• 经验教训总结",
                "总结": f"• {topic}的主要收获\n• 关键要点回顾\n• 下一步行动计划"
            }
            return content_templates.get(section, f"• 关于{section}的要点\n• 详细说明和分析\n• 相关案例或数据")
        else:
            content_templates = {
                "Introduction": f"• Importance of {topic}\n• Objectives of this presentation\n• Overview of main topics",
                "Main Content": f"• Core aspects of {topic}\n• Key features and benefits\n• Practical applications",
                "Case Study": f"• Success story showcase\n• Implementation process\n• Lessons learned",
                "Conclusion": f"• Key takeaways from {topic}\n• Summary of main points\n• Next steps and action plan"
            }
            return content_templates.get(section, f"• Key points about {section}\n• Detailed explanation and analysis\n• Related examples or data")

    def _generate_css_styles(self, config: Dict[str, Any]) -> str:
        """Generate CSS styles for the presentation"""
        return f"""
        * {{
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }}

        body {{
            font-family: {config['font_family']};
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            min-height: 100vh;
            display: flex;
            align-items: center;
            justify-content: center;
        }}

        .presentation-container {{
            width: 1280px;
            height: 720px;
            background: white;
            border-radius: 15px;
            box-shadow: 0 20px 40px rgba(0,0,0,0.1);
            overflow: hidden;
            margin: 0 auto;
        }}

        .slides-wrapper {{
            position: relative;
            width: 100%;
            height: 100%;
            overflow: hidden;
        }}

        .slide {{
            display: none;
            padding: 80px;
            width: 100%;
            height: 100%;
            background: white;
            position: relative;
            box-sizing: border-box;
        }}

        .slide.active {{
            display: block;
        }}

        .slide h1 {{
            color: {config['color_scheme']};
            font-size: 2.5em;
            margin-bottom: 20px;
            text-align: center;
            font-weight: bold;
        }}

        .slide h2 {{
            color: {config['color_scheme']};
            font-size: 2em;
            margin-bottom: 30px;
            border-bottom: 3px solid {config['color_scheme']};
            padding-bottom: 10px;
        }}

        .slide h3 {{
            color: #555;
            font-size: 1.2em;
            margin-bottom: 20px;
            text-align: center;
            font-style: italic;
        }}

        .slide .content {{
            font-size: 1.2em;
            line-height: 1.8;
            color: #333;
            white-space: pre-line;
        }}

        .slide.title-slide {{
            text-align: center;
            display: flex;
            flex-direction: column;
            justify-content: center;
            align-items: center;
            background: linear-gradient(135deg, {config['color_scheme']}22, {config['color_scheme']}11);
        }}

        .slide.title-slide h1 {{
            font-size: 3.5em;
            margin-bottom: 30px;
            color: {config['color_scheme']};
        }}

        .slide.title-slide h3 {{
            font-size: 1.5em;
            color: #666;
        }}

        .slide.agenda-slide ul {{
            list-style: none;
            padding: 0;
        }}

        .slide.agenda-slide li {{
            padding: 15px 0;
            font-size: 1.3em;
            border-bottom: 1px solid #eee;
            color: #555;
        }}

        .slide.content-slide .content {{
            margin-top: 20px;
        }}

        .slide.thankyou-slide {{
            text-align: center;
            display: flex;
            flex-direction: column;
            justify-content: center;
            align-items: center;
            background: linear-gradient(135deg, {config['color_scheme']}22, {config['color_scheme']}11);
        }}

        .navigation {{
            background: {config['color_scheme']};
            padding: 20px;
            display: flex;
            justify-content: space-between;
            align-items: center;
            color: white;
        }}

        .navigation button {{
            background: rgba(255,255,255,0.2);
            border: none;
            color: white;
            padding: 10px 20px;
            border-radius: 5px;
            cursor: pointer;
            font-size: 1em;
            transition: background 0.3s;
        }}

        .navigation button:hover {{
            background: rgba(255,255,255,0.3);
        }}

        .navigation button:disabled {{
            opacity: 0.5;
            cursor: not-allowed;
        }}

        #slideCounter {{
            font-weight: bold;
            font-size: 1.1em;
        }}

        /* Responsive design for 1280x720 slides */
        @media (max-width: 1280px) {{
            .presentation-container {{
                width: 100vw;
                height: 56.25vw;
                max-height: 100vh;
                transform: none;
            }}
        }}

        @media (max-width: 1024px) {{
            .presentation-container {{
                transform: scale(0.8);
                transform-origin: center;
            }}
        }}

        @media (max-width: 768px) {{
            .presentation-container {{
                transform: scale(0.6);
                transform-origin: center;
            }}

            .slide {{
                padding: 40px 30px;
            }}

            .slide h1 {{
                font-size: 2em;
            }}

            .slide h2 {{
                font-size: 1.8em;
            }}

            .slide .content {{
                font-size: 1.1em;
            }}

            .navigation {{
                padding: 15px;
            }}

            .navigation button {{
                padding: 10px 15px;
                font-size: 0.9em;
            }}
        }}

        @media (max-width: 480px) {{
            .presentation-container {{
                transform: scale(0.4);
                transform-origin: center;
            }}
        }}
        """

    def _generate_slide_html(self, slide: Dict[str, Any], config: Dict[str, Any]) -> str:
        """Generate HTML for a single slide"""
        slide_type = slide.get("type", "content")
        slide_id = slide.get("id", 1)
        title = slide.get("title", "")
        subtitle = slide.get("subtitle", "")
        content = slide.get("content", "")

        if slide_type == "title":
            return f"""
            <div class="slide title-slide" id="slide-{slide_id}">
                <h1>{title}</h1>
                {f'<h3>{subtitle}</h3>' if subtitle else ''}
            </div>
            """

        elif slide_type == "agenda":
            # Convert content to HTML list
            agenda_items = [item.strip() for item in content.split('\n') if item.strip()]
            agenda_html = '<ul>' + ''.join([f'<li>{item}</li>' for item in agenda_items]) + '</ul>'

            return f"""
            <div class="slide agenda-slide" id="slide-{slide_id}">
                <h2>{title}</h2>
                {agenda_html}
            </div>
            """

        elif slide_type == "thankyou":
            return f"""
            <div class="slide thankyou-slide" id="slide-{slide_id}">
                <h1>{title}</h1>
                {f'<h3>{subtitle}</h3>' if subtitle else ''}
            </div>
            """

        else:  # content slide
            return f"""
            <div class="slide content-slide" id="slide-{slide_id}">
                <h2>{title}</h2>
                {f'<h3>{subtitle}</h3>' if subtitle else ''}
                <div class="content">{content}</div>
            </div>
            """
